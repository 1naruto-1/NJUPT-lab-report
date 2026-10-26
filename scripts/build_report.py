"""把 Markdown 报告内容套进官方模板，生成 docx（可选导出 pdf 和逐页预览图）。

用法：
    python build_report.py report.md 输出.docx --project "实验项目第一行" [--project "第二行"]
                           [--kind prelab|report] [--pdf] [--preview 预览目录]

Markdown 约定（只支持下面这些，写别的会原样当正文）：
    # 实验名称                      一级标题（居中）
    ## 实验目的                     二级标题，自动编号 一、二、三……（不要自己写序号）
    ### 1. 实验内容一               小标题（加粗，不缩进）
    普通段落                        首行缩进两字；空行分段
    **加粗**  $行内公式$  H~2~O  x^2^
    $$ 独立公式 $$                  单独成行，居中
    - 条目                          无序条目
    ![图1 说明文字](图片路径){width=12cm}   图片居中，图题放在图下方
    表1 表格名称                    紧接着写管道表格，表题放在表上方
    | 表头 | 表头 |
    | --- | --- |
    | 数据 | 数据 |
    <pagebreak>                     分页
"""
from __future__ import annotations

import argparse
import copy
import io
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
from PIL import Image, ImageOps

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "assets" / "report_template.docx"
CONFIG_PATH = HERE.parent / "config.json"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
BODY_PT, SMALL_PT = 12, 10.5          # 小四正文；五号用于图表题和表格
MAX_W, MAX_H = 15.0, 11.0             # 图片最大宽高（cm）

INLINE = re.compile(r"(\*\*.+?\*\*|\$[^$\n]+?\$|~[^~\s]+~|\^[^^\s]+\^)")


# ---------------------------------------------------------------------------
# 公式：一次调用 pandoc，把全部 LaTeX 片段转成 Word 原生公式（OMML）
# ---------------------------------------------------------------------------
def latex_to_omml(snippets: list[tuple[str, bool]]) -> list[etree._Element]:
    if not snippets:
        return []
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "m.md"
        out = Path(tmp) / "m.docx"
        md.write_text("\n\n".join(f"$${s}$$" if display else f"${s}$" for s, display in snippets),
                      encoding="utf-8")
        subprocess.run(["pandoc", str(md), "-o", str(out)], check=True)
        xml = zipfile.ZipFile(out).read("word/document.xml")
    root = etree.fromstring(xml)
    result = []
    for p, (_, display) in zip(root.iter(f"{{{W_NS}}}p"), snippets):
        tag = "oMathPara" if display else "oMath"
        found = p.find(f".//{{{M_NS}}}{tag}")
        if found is None:
            raise ValueError("pandoc 未能转换公式：" + _)
        result.append(found)
    return result


# ---------------------------------------------------------------------------
# 段落与文字
# ---------------------------------------------------------------------------
def set_run_font(run, size=BODY_PT, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, spacing=1.25):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.alignment = align
    fmt.line_spacing = spacing
    fmt.space_before = fmt.space_after = Pt(0)
    if indent:
        ind = p._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:firstLineChars"), "200")
    return p


class Writer:
    def __init__(self, doc, math_queue):
        self.doc = doc
        self.math = math_queue

    def inline(self, p, text, size=BODY_PT, bold=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                self.inline(p, part[2:-2], size, True)
            elif part.startswith("$") and part.endswith("$") and len(part) > 2:
                p._p.append(copy.deepcopy(self.math.pop(0)))
            elif part.startswith("~") and part.endswith("~") and len(part) > 2:
                run = p.add_run(part[1:-1])
                set_run_font(run, size, bold)
                run.font.subscript = True
            elif part.startswith("^") and part.endswith("^") and len(part) > 2:
                run = p.add_run(part[1:-1])
                set_run_font(run, size, bold)
                run.font.superscript = True
            else:
                set_run_font(p.add_run(part), size, bold)


def heading(doc, text, level):
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    p.paragraph_format.line_spacing = 1.0
    if level == 1:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 22, True)
    else:
        # 与模板一致：编号列表 numId=3 生成 一、二、三……
        pPr = p._p.get_or_add_pPr()
        num = OxmlElement("w:numPr")
        for tag, val in (("w:ilvl", "0"), ("w:numId", "3")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), val)
            num.append(el)
        pPr.insert(1, num)
        run = p.add_run(re.sub(r"^[一二三四五六七八九十]+[、.．]\s*", "", text))
        set_run_font(run, 14, True)
    run.font.color.rgb = None


def picture(doc, path: Path, width_cm: float | None):
    img = ImageOps.exif_transpose(Image.open(path))   # 手机照片按 EXIF 方向摆正
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    w_px, h_px = img.size
    width = min(width_cm or MAX_W, MAX_W)
    if width * h_px / w_px > MAX_H:
        width = MAX_H * w_px / h_px
    buf = io.BytesIO()
    img.save(buf, format="PNG" if path.suffix.lower() == ".png" else "JPEG", quality=92)
    buf.seek(0)
    p = new_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(buf, width=Cm(width))


def caption(doc, writer, text, keep_next=False):
    p = new_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = keep_next
    writer.inline(p, text, SMALL_PT)


def table(doc, writer, rows: list[list[str]]):
    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            writer.inline(p, text.strip(), SMALL_PT, bold=(i == 0))
            # 除最后一行外都“与下段同页”，整张表连同表题尽量放在同一页
            p.paragraph_format.keep_with_next = i < len(rows) - 1
        trPr = t.rows[i]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    # 表后留一点空
    new_paragraph(doc, indent=False, spacing=1.0)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
            continue
        rows.append(cells)
    return rows


# ---------------------------------------------------------------------------
# 封面
# ---------------------------------------------------------------------------
def fill_cover(doc, projects: list[str], kind: str):
    if not CONFIG_PATH.exists():
        raise SystemExit(f"缺少 {CONFIG_PATH}：复制 config.example.json 为 config.json 并填写学生信息")
    info = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))["student"]
    labels = {"院：": info["college"], "级：": info["class"], "号：": info["id"],
              "名：": info["name"], "指导教师：": info["teacher"], "课程名称：": info["course"]}
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        text = p.text
        if kind == "prelab" and "实验报告" in text and p.runs and p.runs[-1].font.size and p.runs[-1].font.size.pt > 30:
            for run in p.runs:
                if run.text == "实验报告":
                    run.text = "实验预习报告"
                # 多了两个字，42 磅会折行，整行统一缩到 38 磅
                if run.text and run.font.size and run.font.size.pt > 30:
                    run.font.size = Pt(38)
        if text.strip().startswith("实验") and "项目" in text:
            underlined = [r for r in p.runs if r.underline]
            lines = [paras[i], paras[i + 1]]
            for k, line in enumerate(lines):
                runs = [r for r in line.runs if r.underline]
                if not runs:
                    continue
                runs[0].text = projects[k] if k < len(projects) else ""
                for r in runs[1:]:
                    r.text = ""
            continue
        for label, value in labels.items():
            if label in text:
                runs = [r for r in p.runs if r.underline]
                if runs:
                    runs[0].text = value
                    for r in runs[1:]:
                        r.text = ""
        if "学年第" in text:
            runs = [r for r in p.runs if r.underline]
            if len(runs) >= 2:
                runs[0].text, runs[1].text = info["term_year"], info["term_no"]


def clear_body(doc):
    """删除模板中从“实验名称”标题开始的示例正文，保留封面、页眉和分节设置。"""
    body = doc.element.body
    started = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        style = child.find(f".//{{{W_NS}}}pStyle")
        if not started and child.tag == qn("w:p") and style is not None and style.get(qn("w:val")) == "1":
            started = True
        if started:
            body.remove(child)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def build(md_text: str, out: Path, projects: list[str], kind: str, base: Path):
    blocks = md_text.replace("\r\n", "\n").split("\n")
    # 先收集所有公式，一次性转换
    snippets = []
    for line in blocks:
        s = line.strip()
        if s.startswith("$$") and s.endswith("$$") and len(s) > 4:
            snippets.append((s[2:-2].strip(), True))
        else:
            snippets += [(m[1:-1], False) for m in re.findall(r"\$[^$\n]+?\$", s)]
    math = latex_to_omml(snippets)

    doc = Document(str(TEMPLATE))
    fill_cover(doc, projects, kind)
    clear_body(doc)
    w = Writer(doc, math)

    i, para_buf = 0, []

    def flush():
        if para_buf:
            p = new_paragraph(doc)
            w.inline(p, "".join(para_buf))
            para_buf.clear()

    while i < len(blocks):
        line = blocks[i].rstrip()
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("### "):
            flush()
            p = new_paragraph(doc, indent=False)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.keep_with_next = True
            w.inline(p, s[4:], BODY_PT, True)
        elif s.startswith("## "):
            flush()
            heading(doc, s[3:].strip(), 2)
        elif s.startswith("# "):
            flush()
            heading(doc, s[2:].strip(), 1)
        elif s.startswith("$$") and s.endswith("$$") and len(s) > 4:
            flush()
            p = new_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False)
            p._p.append(copy.deepcopy(w.math.pop(0)))
        elif s == "<pagebreak>":
            flush()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif m := re.fullmatch(r"!\[(.*?)\]\((.+?)\)(\{width=([\d.]+)cm\})?", s):
            flush()
            path = (base / m.group(2)).resolve() if not Path(m.group(2)).is_absolute() else Path(m.group(2))
            picture(doc, path, float(m.group(4)) if m.group(4) else None)
            caption(doc, w, m.group(1))
        elif re.match(r"^表\s*\d+", s) and i + 1 < len(blocks) and blocks[i + 1].strip().startswith("|"):
            flush()
            caption(doc, w, s, keep_next=True)
            j = i + 1
            while j < len(blocks) and blocks[j].strip().startswith("|"):
                j += 1
            table(doc, w, parse_table(blocks[i + 1:j]))
            i = j
            continue
        elif s.startswith("|"):
            flush()
            j = i
            while j < len(blocks) and blocks[j].strip().startswith("|"):
                j += 1
            table(doc, w, parse_table(blocks[i:j]))
            i = j
            continue
        elif s.startswith("- "):
            flush()
            p = new_paragraph(doc, indent=False)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-12)
            w.inline(p, "• " + s[2:])
        else:
            para_buf.append(s)
        i += 1
    flush()
    if w.math:
        raise ValueError("公式数量对不上，检查 $ 是否成对")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def export_pdf(docx: Path) -> Path:
    pdf = docx.with_suffix(".pdf")
    try:
        import win32com.client
        word = win32com.client.DispatchEx("Word.Application")   # 独立进程，不影响已打开的 Word
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            d = word.Documents.Open(str(docx), ReadOnly=True)
            d.ExportAsFixedFormat(str(pdf), 17)
            d.Close(False)
        finally:
            word.Quit()
    except Exception:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=True, capture_output=True)
    return pdf


def preview(pdf: Path, outdir: Path) -> list[str]:
    import fitz
    outdir.mkdir(parents=True, exist_ok=True)
    files = []
    for n, page in enumerate(fitz.open(str(pdf)), start=1):
        path = outdir / f"page_{n:02d}.png"
        page.get_pixmap(dpi=80).save(str(path))
        files.append(str(path))
    return files


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("output")
    ap.add_argument("--project", action="append", default=[], help="封面“实验项目”，最多两行")
    ap.add_argument("--kind", choices=["prelab", "report"], default="report")
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--preview", help="把 pdf 逐页渲染成 png 的目录，便于检查排版")
    args = ap.parse_args()
    md = Path(args.markdown).resolve()
    out = Path(args.output).resolve()
    build(md.read_text(encoding="utf-8"), out, args.project, args.kind, md.parent)
    result = {"docx": str(out)}
    if args.pdf or args.preview:
        pdf = export_pdf(out)
        result["pdf"] = str(pdf)
        if args.preview:
            result["pages"] = preview(pdf, Path(args.preview))
    print(json.dumps(result, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
